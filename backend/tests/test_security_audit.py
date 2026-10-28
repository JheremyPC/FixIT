"""Security hardening, audit trail (query/export/integrity), terms acceptance and removal of the demo access."""
import base64
import csv
import io
import json
import uuid
from datetime import datetime, timedelta, timezone

import jwt
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.config import settings
from app.core.database import SessionLocal
from app.legal import TERMS
from app.main import app
from app.models import AuditLog, TermsAcceptance
from app.security.hardening import login_throttle
from app.services.audit import REDACTED, audit, clean_text, sealed_fields, verify
from app.services.exports import safe_cell

PASSWORD = "Secure!123"
ROLES = ["ADMIN", "JEFE_TI", "TECNICO", "USUARIO"]
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32


def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def fresh_throttle():
    login_throttle.clear()
    yield
    login_throttle.clear()


def login(client, email, password=PASSWORD):
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


@pytest.fixture(scope="module")
def admin(client):
    return auth(login(client, "admin@fixit.org", "FixIT!2026"))


@pytest.fixture(scope="module")
def accounts(client, admin):
    """One account per role (plus a second requester) with real sessions, and a ticket of the first requester."""
    suffix = uuid.uuid4().hex[:8]
    out = {"ADMIN": {"headers": admin}}
    specialty = client.post("/api/catalogs/specialties", headers=admin, json={"name": f"Seguridad {suffix}"}).json()["id"]
    category = client.post("/api/catalogs/categories", headers=admin, json={"name": f"Accesos {suffix}", "specialty_id": specialty}).json()["id"]
    for role, key in [("JEFE_TI", "JEFE_TI"), ("TECNICO", "TECNICO"), ("TECNICO", "TECNICO_2"), ("USUARIO", "USUARIO"), ("USUARIO", "USUARIO_2")]:
        email = f"{key.lower()}-{suffix}@fixit.org"
        user_id = client.post("/api/users", headers=admin, json={"full_name": f"Prueba {key}", "email": email, "password": PASSWORD,
                                                                 "role_code": role}).json()["id"]
        out[key] = {"id": user_id, "email": email, "headers": auth(login(client, email))}
    # Only the first technician has the specialty, so the ticket is assigned to them.
    client.post("/api/technicians", headers=admin, json={"user_id": out["TECNICO"]["id"], "specialty_ids": [specialty]})
    client.post("/api/technicians", headers=admin, json={"user_id": out["TECNICO_2"]["id"], "specialty_ids": [], "available": False})
    area = client.get("/api/catalogs/areas", headers=admin).json()[0]["id"]
    ticket = client.post("/api/tickets", headers=out["USUARIO"]["headers"], json={
        "title": "=HYPERLINK(\"http://example.invalid\",\"x\")", "description": "Prueba de inyección de fórmulas en reportes.",
        "area_id": area, "category_id": category}).json()
    assert ticket["assigned"] is True
    out["ticket"] = ticket["id"]
    return out


def audit_rows(**filters):
    db = SessionLocal()
    try:
        query = db.query(AuditLog)
        for key, value in filters.items():
            query = query.filter(getattr(AuditLog, key) == value)
        return query.order_by(AuditLog.created_at).all()
    finally:
        db.close()


# ---------- Demo access removed ----------

def test_demo_access_is_gone(client):
    assert client.get("/api/auth/demo").status_code == 404
    for role in ROLES:
        assert client.post(f"/api/auth/demo/{role}").status_code == 404
    assert client.get("/static/demo-access.js").status_code == 404
    page = client.get("/").text
    assert "demo-access" not in page and "Accesos de demostración" not in page
    assert not hasattr(settings, "demo_access_enabled")


# ---------- Authentication ----------

def test_normal_login_refresh_and_failures_are_audited(client, accounts):
    email = accounts["USUARIO"]["email"]
    data = client.post("/api/auth/login", json={"email": email, "password": PASSWORD}).json()
    assert data["user"]["role"] == "USUARIO"
    assert client.get("/api/auth/me", headers=auth(data["access_token"])).status_code == 200
    assert client.post("/api/auth/refresh", json={"refresh_token": data["refresh_token"]}).status_code == 200
    # Same message for a wrong password and for an unknown account.
    wrong = client.post("/api/auth/login", json={"email": email, "password": "Incorrecta1"})
    unknown = client.post("/api/auth/login", json={"email": f"nadie-{uuid.uuid4().hex[:6]}@fixit.org", "password": "Incorrecta1"})
    assert wrong.status_code == unknown.status_code == 401 and wrong.json() == unknown.json()
    failed = audit_rows(action="LOGIN_FALLIDO", user_id=accounts["USUARIO"]["id"])
    assert failed and failed[-1].outcome == "FALLO" and verify(failed[-1]) == "VALIDO"
    assert "Incorrecta1" not in json.dumps(failed[-1].new_value)


def test_login_is_throttled_per_account(client, accounts, monkeypatch):
    monkeypatch.setattr(settings, "login_max_failures", 3)
    email = accounts["TECNICO_2"]["email"]
    for _ in range(3):
        assert client.post("/api/auth/login", json={"email": email, "password": "Incorrecta1"}).status_code == 401
    blocked = client.post("/api/auth/login", json={"email": email, "password": PASSWORD})
    assert blocked.status_code == 429 and int(blocked.headers["Retry-After"]) > 0
    assert audit_rows(action="LOGIN_BLOQUEADO", user_id=accounts["TECNICO_2"]["id"])[-1].outcome == "DENEGADO"
    # Other accounts are not affected, and the block ends with the window.
    assert client.post("/api/auth/login", json={"email": accounts["TECNICO"]["email"], "password": PASSWORD}).status_code == 200
    login_throttle.clear()
    assert client.post("/api/auth/login", json={"email": email, "password": PASSWORD}).status_code == 200


def test_tokens_are_strictly_validated(client, accounts):
    data = client.post("/api/auth/login", json={"email": accounts["JEFE_TI"]["email"], "password": PASSWORD}).json()
    claims = jwt.decode(data["access_token"], options={"verify_signature": False})
    now = datetime.now(timezone.utc)
    forged = [
        jwt.encode({**claims, "role": "ADMIN"}, "otra-clave-que-no-es-la-del-servidor-0123456789", algorithm="HS256"),
        jwt.encode({**claims, "exp": now - timedelta(minutes=1)}, settings.jwt_secret_key, algorithm="HS256"),
        unsigned_token({**claims, "role": "ADMIN"}),
        data["refresh_token"],  # a refresh token is not an access token
        "no-es-un-jwt",
    ]
    for token in forged:
        assert client.get("/api/auth/me", headers=auth(token)).status_code == 401
    assert client.post("/api/auth/refresh", json={"refresh_token": data["access_token"]}).status_code == 401
    # Deactivated accounts lose their open sessions immediately.
    admin, email = accounts["ADMIN"]["headers"], f"tmp-{uuid.uuid4().hex[:6]}@fixit.org"
    user = client.post("/api/users", headers=admin, json={"full_name": "Temporal", "email": email, "password": PASSWORD}).json()["id"]
    session = client.post("/api/auth/login", json={"email": email, "password": PASSWORD}).json()
    assert client.patch(f"/api/users/{user}", headers=admin, json={"active": False}).status_code == 200
    assert client.get("/api/auth/me", headers=auth(session["access_token"])).status_code == 401
    assert client.post("/api/auth/refresh", json={"refresh_token": session["refresh_token"]}).status_code == 401
    assert client.post("/api/auth/login", json={"email": email, "password": PASSWORD}).status_code == 401


def unsigned_token(claims: dict) -> str:
    """JWT with "alg": "none" (no signature), the classic algorithm-confusion attempt."""
    part = lambda obj: base64.urlsafe_b64encode(json.dumps(obj, default=str).encode()).rstrip(b"=").decode()
    return f"{part({'alg': 'none', 'typ': 'JWT'})}.{part(claims)}."


def test_validation_errors_do_not_echo_input(client):
    response = client.post("/api/auth/login", json={"email": "admin@fixit.org", "password": "corta"})
    assert response.status_code == 422
    assert "corta" not in response.text and all(set(e) == {"loc", "msg", "type"} for e in response.json()["errors"])


def test_admin_cannot_lock_themselves_out(client, admin):
    me = client.get("/api/auth/me", headers=admin).json()
    assert client.patch(f"/api/users/{me['id']}", headers=admin, json={"active": False}).status_code == 422
    assert client.patch(f"/api/users/{me['id']}", headers=admin, json={"role_code": "USUARIO"}).status_code == 422
    assert client.get("/api/auth/me", headers=admin).json()["role"] == "ADMIN"


# ---------- Authorization and object-level access ----------

def test_role_matrix(client, accounts):
    checks = [("/api/users", {"ADMIN"}), ("/api/settings", {"ADMIN"}), ("/api/technicians", {"ADMIN", "JEFE_TI"}),
              ("/api/audit", {"ADMIN", "JEFE_TI"}), ("/api/audit/facets", {"ADMIN", "JEFE_TI"}), ("/api/audit/export", {"ADMIN", "JEFE_TI"}),
              ("/api/reports/tickets.csv", {"ADMIN", "JEFE_TI"}), ("/api/ml/status", {"ADMIN", "JEFE_TI"}), ("/api/agents/monitor", {"ADMIN", "JEFE_TI"})]
    for role in ROLES:
        headers = accounts[role]["headers"]
        for path, allowed in checks:
            assert client.get(path, headers=headers).status_code == (200 if role in allowed else 403), (role, path)
    for path, _ in checks:
        assert client.get(path).status_code == 401, path


def test_denied_audit_access_is_recorded(client, accounts):
    before = len(audit_rows(action="ACCESO_DENEGADO", user_id=accounts["TECNICO"]["id"]))
    assert client.get("/api/audit/export?format=json", headers=accounts["TECNICO"]["headers"]).status_code == 403
    rows = audit_rows(action="ACCESO_DENEGADO", user_id=accounts["TECNICO"]["id"])
    assert len(rows) == before + 1 and rows[-1].outcome == "DENEGADO" and rows[-1].new_value["path"] == "/api/audit/export"


def test_audit_records_cannot_be_changed_through_the_api(client, admin):
    record_id = client.get("/api/audit?limit=1", headers=admin).json()[0]["id"]
    for method in ("put", "patch", "delete"):
        assert getattr(client, method)(f"/api/audit/{record_id}", headers=admin).status_code in {404, 405}
        assert getattr(client, method)("/api/audit", headers=admin).status_code == 405


def test_tickets_and_attachments_are_isolated(client, accounts):
    ticket = accounts["ticket"]
    owner, other = accounts["USUARIO"]["headers"], accounts["USUARIO_2"]["headers"]
    upload = client.post(f"/api/tickets/{ticket}/attachments", headers=owner, files={"file": ("../../captura.png", PNG, "image/png")})
    assert upload.status_code == 201
    attachment = upload.json()["id"]
    for headers, expected in [(owner, 200), (accounts["TECNICO"]["headers"], 200), (accounts["JEFE_TI"]["headers"], 200),
                              (other, 403), (accounts["TECNICO_2"]["headers"], 403)]:
        assert client.get(f"/api/tickets/{ticket}", headers=headers).status_code == expected
        assert client.get(f"/api/attachments/{attachment}", headers=headers).status_code == expected
    assert client.post(f"/api/tickets/{ticket}/comments", headers=other, json={"body": "hola"}).status_code == 403
    assert client.post(f"/api/llm/tickets/{ticket}/summary", headers=other).status_code == 403
    assert all(t["id"] != ticket for t in client.get("/api/tickets", headers=other).json())
    detail = client.get(f"/api/tickets/{ticket}", headers=owner).json()
    assert detail["attachments"][0]["name"] == "captura.png"
    download = client.get(f"/api/attachments/{attachment}", headers=owner)
    assert download.headers["content-disposition"].startswith("attachment") and download.headers["x-content-type-options"] == "nosniff"


def test_attachment_content_must_match_its_type(client, accounts):
    ticket, owner = accounts["ticket"], accounts["USUARIO"]["headers"]
    url = f"/api/tickets/{ticket}/attachments"
    assert client.post(url, headers=owner, files={"file": ("x.png", b"<script>alert(1)</script>", "image/png")}).status_code == 415
    assert client.post(url, headers=owner, files={"file": ("x.html", b"<html></html>", "text/html")}).status_code == 415
    assert client.post(url, headers=owner, files={"file": ("x.txt", b"\xff\xfe\x00binario", "text/plain")}).status_code == 415
    assert client.post(url, headers=owner, files={"file": ("x.pdf", b"", "application/pdf")}).status_code == 413
    assert client.post(url, headers=owner, files={"file": ("nota.txt", "Texto válido".encode(), "text/plain")}).status_code == 201


# ---------- Audit query, export and integrity ----------

def test_audit_list_filters_and_integrity(client, admin, accounts):
    rows = client.get(f"/api/audit?action=LOGIN&user_id={accounts['USUARIO']['id']}", headers=admin)
    assert rows.status_code == 200 and int(rows.headers["X-Total-Count"]) >= 1
    assert all(r["action"] == "LOGIN" and r["user_id"] == accounts["USUARIO"]["id"] and r["integrity"] == "VALIDO" for r in rows.json())
    assert all(r["created_at"].endswith("Z") for r in rows.json())
    future = (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()
    assert client.get("/api/audit", headers=admin, params={"date_from": future}).json() == []
    assert client.get("/api/audit", headers=admin, params={"date_from": future, "date_to": future}).status_code == 422
    assert client.get("/api/audit?outcome=QUIZAS", headers=admin).status_code == 422
    assert client.get("/api/audit?action=LOGIN';DROP TABLE audit_logs;--", headers=admin).status_code == 422
    assert client.get("/api/audit", headers=admin, params={"q": "%_\\"}).status_code == 200  # LIKE wildcards are escaped
    system = client.get("/api/audit?user_id=SISTEMA&limit=5", headers=admin).json()
    assert all(r["user_id"] is None for r in system)
    facets = client.get("/api/audit/facets", headers=admin).json()
    assert "LOGIN" in facets["actions"] and set(facets["outcomes"]) == {"EXITO", "FALLO", "DENEGADO"}
    assert audit_rows(action="CONSULTAR")[-1].module == "AUDITORIA"


def test_csv_export_is_faithful_safe_and_audited(client, admin, accounts):
    email = f"formula-{uuid.uuid4().hex[:6]}@fixit.org"
    client.post("/api/users", headers=admin, json={"full_name": "=1+2 Formula", "email": email, "password": PASSWORD})
    user_id = next(u["id"] for u in client.get("/api/users", headers=admin).json() if u["email"] == email)
    login(client, email)
    response = client.get(f"/api/audit/export?format=csv&user_id={user_id}", headers=admin)
    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")
    disposition = response.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="fixit-auditoria_') and disposition.endswith('.csv"')
    body = response.content.decode("utf-8")
    assert body.startswith("﻿")
    lines = list(csv.reader(io.StringIO(body.lstrip("﻿"))))
    meta_rows = int(lines[0][1])
    meta = {row[0]: row[1] for row in lines[:meta_rows]}
    assert meta["# zona_horaria"] == "UTC" and admin_email(client, admin) in meta["# exportado_por"]
    assert json.loads(meta["# filtros"]) == {"usuario_id": user_id}
    header, records = lines[meta_rows + 1], lines[meta_rows + 2:]
    assert header[:3] == ["id", "fecha_utc", "usuario_id"] and "integridad" in header
    assert {r[header.index("accion")] for r in records} == {"LOGIN"}
    assert all(r[header.index("usuario_nombre")] == "'=1+2 Formula" for r in records)  # neutralized formula
    assert all(r[header.index("integridad")] == "VALIDO" for r in records)
    assert PASSWORD not in body and "password_hash" not in body
    event = audit_rows(action="EXPORTAR")[-1]
    assert event.user_id is not None and event.new_value["formato"] == "csv" and event.new_value["registros"] == len(records)
    assert meta["# evento_exportacion_id"] == event.id


def admin_email(client, admin):
    return client.get("/api/auth/me", headers=admin).json()["email"]


def test_json_export_metadata_and_limits(client, admin, accounts, monkeypatch):
    start = (datetime.now(timezone.utc) - timedelta(hours=2)).isoformat()
    response = client.get("/api/audit/export", headers=admin, params={"format": "json", "module": "AUTH", "date_from": start})
    assert response.status_code == 200 and response.headers["content-type"].startswith("application/json")
    data = response.json()
    meta, records = data["metadatos"], data["registros"]
    assert meta["zona_horaria"] == "UTC" and meta["exportado_en"].endswith("Z") and meta["total_registros"] == len(records)
    assert meta["exportado_por"]["rol"] == "ADMIN" and meta["aplicacion"]["nombre"] == "FixIT" and meta["aplicacion"]["version"]
    assert meta["filtros"]["modulo"] == "AUTH" and meta["intervalo_solicitado"]["desde"].startswith(start[:16])
    assert records and all(r["modulo"] == "AUTH" for r in records)
    assert [r["fecha_utc"] for r in records] == sorted(r["fecha_utc"] for r in records)
    assert meta["intervalo_registros"]["primero"] == records[0]["fecha_utc"]
    assert sum(meta["integridad"]["resumen"].values()) == len(records)
    monkeypatch.setattr(settings, "audit_export_max_rows", 1)
    too_many = client.get("/api/audit/export?format=json", headers=admin)
    assert too_many.status_code == 422 and "Acote" in too_many.json()["detail"]
    assert client.get("/api/audit/export?format=xml", headers=admin).status_code == 422


def test_seal_detects_tampering_and_secrets_are_redacted(client, admin):
    db = SessionLocal()
    try:
        row = audit(db, None, "PRUEBA\r\nLOGIN", "PRUEBA", "TEST", None,
                    new={"password": "secreta", "nested": {"refresh_token": "abc", "api_key": "k"}, "ok": 1})
        unsealed = AuditLog(**{k: v for k, v in sealed_fields(None, "PRUEBA_SIN_SELLO", "PRUEBA", "TEST", None).items() if k != "integrity_hash"})
        db.add(unsealed)
        db.commit()
        row_id, unsealed_id = row.id, unsealed.id
        assert row.action == "PRUEBA\\x0d\\x0aLOGIN"  # CR/LF cannot forge a second line
        assert row.new_value == {"password": REDACTED, "nested": {"refresh_token": REDACTED, "api_key": REDACTED}, "ok": 1}
        # Direct database edit (SQLite has no append-only trigger; PostgreSQL rejects this UPDATE).
        db.execute(text("UPDATE audit_logs SET entity = 'MANIPULADO' WHERE id = :id"), {"id": row_id})
        db.commit()
    finally:
        db.close()
    by_id = {r["id"]: r for r in client.get("/api/audit?module=PRUEBA&limit=50", headers=admin).json()}
    assert by_id[row_id]["integrity"] == "ALTERADO" and by_id[unsealed_id]["integrity"] == "SIN_SELLO"
    assert "secreta" not in json.dumps(by_id[row_id])


def test_spreadsheet_injection_is_neutralized(client, admin, accounts):
    assert safe_cell("=SUM(A1)") == "'=SUM(A1)" and safe_cell("  @cmd") == "'  @cmd" and safe_cell("-1+2") == "'-1+2"
    assert safe_cell(-5) == -5 and safe_cell("Normal") == "Normal" and safe_cell(None) == ""
    assert clean_text("a\x1b[31mb c") == "a\\x1b[31mb\\u2028c"
    report = client.get("/api/reports/tickets.csv", headers=admin).text
    assert '\'=HYPERLINK(""http://example.invalid"",""x"")' in report
    sheet = load_workbook(io.BytesIO(client.get("/api/reports/tickets.xlsx", headers=admin).content)).active
    cells = [c for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=HYPERLINK")]
    assert cells and all(c.data_type == "s" for c in cells)


# ---------- Terms and privacy ----------

def test_legal_documents_are_public_and_flag_pending_data(client):
    data = client.get("/api/legal").json()
    ids = {d["id"]: d for d in data["documents"]}
    assert set(ids) == {"terminos", "privacidad"}
    assert ids["terminos"]["requires_acceptance"] is True and ids["privacidad"]["requires_acceptance"] is False
    assert data["pending_marker"] in json.dumps(ids["privacidad"], ensure_ascii=False)
    assert "29733" in json.dumps(ids["privacidad"], ensure_ascii=False)
    assert client.get("/legal").status_code == 200


def test_registration_requires_express_acceptance(client):
    email = f"registro-{uuid.uuid4().hex[:6]}@fixit.org"
    base = {"full_name": "Nueva Solicitante", "email": email, "password": PASSWORD, "role_code": "ADMIN"}
    assert client.post("/api/auth/register", json=base).status_code == 422
    assert client.post("/api/auth/register", json={**base, "accept_terms": True, "terms_version": "0.0"}).status_code == 409
    created = client.post("/api/auth/register", json={**base, "accept_terms": True, "terms_version": TERMS["version"]})
    assert created.status_code == 201
    token = login(client, email)
    assert client.get("/api/auth/me", headers=auth(token)).json()["role"] == "USUARIO"  # requested role is ignored
    state = client.get("/api/legal/acceptances", headers=auth(token)).json()
    assert state["pending"] == [] and state["accepted"][0]["version"] == TERMS["version"]
    db = SessionLocal()
    try:
        row = db.query(TermsAcceptance).filter_by(user_id=created.json()["id"]).one()
        assert row.document == "terminos" and row.accepted_at is not None
    finally:
        db.close()


def test_existing_users_accept_terms_explicitly(client, accounts):
    headers = accounts["JEFE_TI"]["headers"]
    assert client.get("/api/legal/acceptances", headers=headers).json()["pending"][0]["document"] == "terminos"
    assert client.post("/api/legal/acceptances", headers=headers, json={"document": "privacidad", "version": "0.1-borrador"}).status_code == 422
    assert client.post("/api/legal/acceptances", headers=headers, json={"document": "terminos", "version": "9.9"}).status_code == 409
    first = client.post("/api/legal/acceptances", headers=headers, json={"document": "terminos", "version": TERMS["version"]})
    again = client.post("/api/legal/acceptances", headers=headers, json={"document": "terminos", "version": TERMS["version"]})
    assert first.status_code == again.status_code == 201 and first.json()["pending"] == [] and len(again.json()["accepted"]) == 1
    assert len(audit_rows(action="ACEPTAR_TERMINOS", user_id=accounts["JEFE_TI"]["id"])) == 1
    assert client.post("/api/legal/acceptances", json={"document": "terminos", "version": TERMS["version"]}).status_code == 401


# ---------- HTTP hardening ----------

def test_security_headers_and_docs_disabled(client, admin):
    for path, headers in [("/", {}), ("/legal", {}), ("/api/auth/me", admin), ("/static/app.js", {})]:
        response = client.get(path, headers=headers)
        assert response.status_code == 200, path
        csp = response.headers["content-security-policy"]
        assert "script-src 'self';" in csp and "frame-ancestors 'none'" in csp and "unsafe-eval" not in csp
        assert "style-src 'self';" in csp and "https:" not in csp  # no third-party origins, no inline <style>
        assert response.headers["cross-origin-embedder-policy"] == "require-corp"
        assert response.headers["x-content-type-options"] == "nosniff" and response.headers["x-frame-options"] == "DENY"
    assert client.get("/api/auth/me", headers=admin).headers["cache-control"] == "no-store"
    assert client.get("/static/app.js").headers["cache-control"] == "no-cache"
    assert client.get("/static/fonts/Geist-Variable.woff2").content[:4] == b"wOF2"
    assert "strict-transport-security" in client.get("/health", headers={"X-Forwarded-Proto": "https"}).headers
    for path in ("/docs", "/redoc", "/openapi.json"):
        assert client.get(path).status_code == 404


# ---------- Regressions found by the OWASP ZAP active scan ----------

def test_nul_bytes_and_duplicates_are_client_errors(client, admin, accounts):
    owner = accounts["USUARIO"]["headers"]
    nul = "\x00"
    assert client.post("/api/agents/triage/suggest", headers=owner, json={"title": f"Red{nul}caida", "description": "x"}).status_code == 422
    assert client.post("/api/auth/login", json={"email": "admin@fixit.org", "password": f"FixIT!2026{nul}"}).status_code == 422
    assert client.post("/api/technicians", headers=admin, json={"user_id": "x", "specialty_ids": [f"a{nul}"]}).status_code == 422
    assert client.get("/api/tickets/%00", headers=admin).status_code in {404, 422}
    level = 7 + int(uuid.uuid4().int % 3)
    code = f"P{uuid.uuid4().hex[:6].upper()}"
    first = client.post("/api/priorities", headers=admin, json={"code": code, "name": "Prueba", "level": level, "sla_hours": 8})
    duplicate = client.post("/api/priorities", headers=admin, json={"code": code, "name": "Prueba", "level": level, "sla_hours": 8})
    assert first.status_code in {201, 409} and duplicate.status_code == 409
